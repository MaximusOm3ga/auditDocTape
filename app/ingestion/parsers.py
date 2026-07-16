import pdfplumber
from docx import Document as DocxDocument
import openpyxl

def parse_pdf(path: str) -> str:

    text = []

    with pdfplumber.open(path) as pdf:

        for page in pdf.pages:

            text.append(page.extract_text() or "")

    return "\n".join(text)

def parse_docx(path: str) -> str:

    doc = DocxDocument(path)

    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def parse_xlsx(path: str) -> str:

    wb = openpyxl.load_workbook(path, data_only=True)

    lines = []

    for sheet in wb.worksheets:

        for row in sheet.iter_rows(values_only=True):

            cells = [str(c) for c in row if c is not None]

            if cells:

                lines.append(" | ".join(cells))

    return "\n".join(lines)

def parse_document(path: str, ext: str) -> str:

    if ext == "pdf":

        return parse_pdf(path)

    if ext == "docx":

        return parse_docx(path)

    if ext in ("xlsx", "xls"):

        return parse_xlsx(path)

    raise ValueError(f"Unsupported extension: {ext}")
